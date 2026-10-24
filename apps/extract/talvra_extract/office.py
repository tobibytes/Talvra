"""DOCX and PPTX."""

import io
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError

from .markdown_sections import sections_from_markdown
from .model import BadFile, Extracted, Section


def _md_escape_cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ").strip()


def _table_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head, *body = rows
    lines = ["| " + " | ".join(_md_escape_cell(c) for c in head) + " |"]
    lines.append("|" + "---|" * width)
    lines += ["| " + " | ".join(_md_escape_cell(c) for c in r) + " |" for r in body]
    return "\n".join(lines)


def _paragraph_md(p: Paragraph) -> str:
    text = p.text.strip()
    if not text:
        return ""
    style = (p.style.name if p.style is not None else "") or ""
    if style == "Title":
        return f"# {text}"
    if style.startswith("Heading"):
        try:
            level = int(style.split()[-1])
        except ValueError:
            level = 2
        return f"{'#' * min(max(level + 1, 2), 6)} {text}"
    if "List" in style:
        return f"- {text}"
    return text


def extract_docx(data: bytes) -> Extracted:
    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError) as e:
        raise BadFile("Couldn't open this Word document. It may be damaged.") from e
    blocks: list[str] = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            md = _paragraph_md(Paragraph(el, doc))
            if md:
                blocks.append(md)
        elif tag == "tbl":
            t = Table(el, doc)
            rows = [[c.text for c in r.cells] for r in t.rows]
            md = _table_md(rows)
            if md:
                blocks.append(md)
    md = "\n\n".join(blocks)
    title = (doc.core_properties.title or "").strip() or None
    return Extracted(md, sections_from_markdown(md, title or "Document"), title=title)


def extract_pptx(data: bytes) -> Extracted:
    try:
        prs = Presentation(io.BytesIO(data))
    except (PptxPackageNotFoundError, zipfile.BadZipFile, KeyError) as e:
        raise BadFile("Couldn't open this PowerPoint file. It may be damaged.") from e
    sections: list[Section] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        lines: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(f"{'  ' * para.level}- {t}")
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                lines.append(_table_md(rows))
        if slide.has_notes_slide:
            notes = (
                slide.notes_slide.notes_text_frame.text.strip()
                if slide.notes_slide.notes_text_frame
                else ""
            )
            if notes:
                lines.append("")
                lines += [
                    f"> {ln}" if ln else ">" for ln in ["**Speaker notes**", *notes.splitlines()]
                ]
        heading = f"Slide {i}: {title}" if title else f"Slide {i}"
        sections.append(Section(heading, 2, "\n".join(lines).strip(), "slide", i))
    md = "\n\n".join(f"## {s.heading}\n\n{s.text}".rstrip() for s in sections)
    first = (
        sections[0].heading.split(": ", 1)[-1] if sections and ": " in sections[0].heading else None
    )
    return Extracted(
        md,
        [s for s in sections if s.text or ": " in s.heading],
        page_count=len(sections),
        title=first,
    )
