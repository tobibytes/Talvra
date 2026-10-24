import os

os.environ.setdefault("ALLOW_NO_AUTH", "1")

import io

import pytest
from docx import Document
from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches


@pytest.fixture
def pdf_bytes() -> bytes:
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for n, body in enumerate(
        [
            "Arrays store elements contiguously. Access by index is O(1).",
            "Linked lists trade random access for cheap insertion. Each node points to the next.",
            "A stack is last-in, first-out. Push and pop are O(1).",
            "A queue is first-in, first-out. Use it for breadth-first search.",
        ],
        start=1,
    ):
        pdf.add_page()
        pdf.cell(text="COSC 220 Fall 2026")  # running header on every page
        pdf.ln(10)
        pdf.multi_cell(0, 8, text=body)
        pdf.ln(10)
        pdf.cell(text=f"Page {n} of 4")
    pdf.set_title("Lecture 1")
    return bytes(pdf.output())


@pytest.fixture
def docx_bytes() -> bytes:
    doc = Document()
    doc.core_properties.title = "Syllabus"
    doc.add_paragraph("COSC 220 Syllabus", style="Title")
    doc.add_heading("Grading", level=1)
    doc.add_paragraph("Labs are 40% of the grade.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Weight"
    t.cell(1, 0).text, t.cell(1, 1).text = "Labs | quizzes", "40%"
    doc.add_heading("Late policy", level=2)
    doc.add_paragraph("Ten percent off per day.", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def pptx_bytes() -> bytes:
    prs = Presentation()
    for title, bullets, notes in [
        (
            "Big-O notation",
            ["Upper bound on growth", "Drop constants"],
            "Mention the exam question.",
        ),
        ("Arrays", ["Contiguous memory"], ""),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            p = body.add_paragraph()
            p.text = b
            p.level = 1
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    blank = prs.slides.add_slide(prs.slide_layouts[6])
    blank.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text_frame.text = "Questions?"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
